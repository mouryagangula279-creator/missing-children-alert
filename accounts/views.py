import os
from django.conf import settings
from django.http import FileResponse, Http404
from django.shortcuts import render, redirect, get_object_or_404
from django.contrib.auth import authenticate, login, logout
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from django.db.models import Q
from django.http import HttpResponse

from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage

from .forms import MissingChildForm
from .models import MissingChild, Alert


# HOME
def home(request):
    return render(
        request,
        "accounts/home.html"
    )


# LOGIN
def login_view(request):

    if request.user.is_authenticated:
        return redirect("dashboard")

    if request.method == "POST":

        username = request.POST.get("username")
        password = request.POST.get("password")

        user = authenticate(
            request,
            username=username,
            password=password
        )

        if user is not None:

            login(request, user)

            return redirect("dashboard")

        else:

            messages.error(
                request,
                "Invalid username or password."
            )

    return render(
        request,
        "accounts/login.html"
    )


# DASHBOARD
@login_required(login_url="login")
def dashboard(request):

    profile = getattr(
        request.user,
        "userprofile",
        None
    )

    total_cases = MissingChild.objects.count()

    total_alerts = Alert.objects.count()

    unread_alerts = Alert.objects.filter(
        is_read=False
    ).count()

    active_cases = MissingChild.objects.filter(
        status="ACTIVE"
    ).count()

    under_review = MissingChild.objects.filter(
        status="UNDER_REVIEW"
    ).count()

    resolved_cases = MissingChild.objects.filter(
        status="RESOLVED"
    ).count()

    recent_cases = MissingChild.objects.order_by(
        "-created_at"
    )[:5]

    return render(
        request,
        "accounts/dashboard.html",
        {
            "profile": profile,
            "total_cases": total_cases,
            "total_alerts": total_alerts,
            "unread_alerts": unread_alerts,
            "active_cases": active_cases,
            "under_review": under_review,
            "resolved_cases": resolved_cases,
            "recent_cases": recent_cases,
        }
    )


# REGISTER MISSING CHILD
@login_required(login_url="login")
def register_missing_child(request):

    profile = getattr(
        request.user,
        "userprofile",
        None
    )

    # Only Admin, Police and Sachivalayam employees can register cases
    if profile is None or profile.role not in [
        "ADMIN",
        "POLICE",
        "SACHIVALAYAM"
    ]:

        messages.error(
            request,
            "You are not authorized to register a missing-child case."
        )

        return redirect("dashboard")

    if request.method == "POST":

        form = MissingChildForm(
            request.POST,
            request.FILES
        )

        if form.is_valid():

            case = form.save(
                commit=False
            )

            # Generate unique case number
            last_case = MissingChild.objects.order_by(
                "-id"
            ).first()

            if last_case:
                next_number = last_case.id + 1
            else:
                next_number = 1

            case.case_number = (
                f"ST-{next_number:05d}"
            )

            case.reported_by = request.user

            case.save()

            # Create alert
            Alert.objects.create(
                case=case,
                alert_type="MISSING_CHILD",
                message=(
                    f"New missing-child case registered: "
                    f"{case.child_name} "
                    f"(Case Number: {case.case_number}). "
                    f"Last seen at "
                    f"{case.location_last_seen}."
                )
            )

            messages.success(
                request,
                f"Missing-child case "
                f"{case.case_number} "
                f"has been registered successfully."
            )

            return redirect(
                "dashboard"
            )

    else:

        form = MissingChildForm()

    return render(
        request,
        "accounts/register_missing_child.html",
        {
            "form": form
        }
    )


# MISSING CASES
@login_required(login_url="login")
def missing_cases(request):

    cases = MissingChild.objects.all().order_by(
        "-created_at"
    )

    search_query = request.GET.get(
        "search",
        ""
    )

    status = request.GET.get(
        "status",
        ""
    )

    gender = request.GET.get(
        "gender",
        ""
    )

    if search_query:

        cases = cases.filter(
            Q(child_name__icontains=search_query) |
            Q(case_number__icontains=search_query)
        )

    if status:

        cases = cases.filter(
            status=status
        )

    if gender:

        cases = cases.filter(
            gender=gender
        )

    profile = getattr(
        request.user,
        "userprofile",
        None
    )

    return render(
        request,
        "accounts/missing_cases.html",
        {
            "cases": cases,
            "profile": profile,
            "search_query": search_query,
            "selected_status": status,
            "selected_gender": gender,
        }
    )


# CASE DETAIL
@login_required(login_url="login")
def case_detail(request, case_id):

    case = get_object_or_404(
        MissingChild,
        id=case_id
    )

    profile = getattr(
        request.user,
        "userprofile",
        None
    )

    case_alerts = Alert.objects.filter(
        case=case
    ).order_by(
        "-created_at"
    )

    return render(
        request,
        "accounts/case_detail.html",
        {
            "case": case,
            "profile": profile,
            "case_alerts": case_alerts,
        }
    )


# UPDATE CASE STATUS
@login_required(login_url="login")
def update_case_status(request, case_id):

    profile = getattr(
        request.user,
        "userprofile",
        None
    )

    if profile is None or profile.role not in [
        "POLICE",
        "ADMIN"
    ]:

        messages.error(
            request,
            "You are not authorized to update case status."
        )

        return redirect(
            "case_detail",
            case_id=case_id
        )

    case = get_object_or_404(
        MissingChild,
        id=case_id
    )

    if request.method == "POST":

        new_status = request.POST.get(
            "status"
        )

        allowed_statuses = [
            "REPORTED",
            "UNDER_REVIEW",
            "ACTIVE",
            "RESOLVED"
        ]

        if new_status in allowed_statuses:

            old_status = case.status

            if old_status != new_status:

                case.status = new_status
                case.save()

                old_status_display = dict(
                    MissingChild.STATUS_CHOICES
                ).get(
                    old_status,
                    old_status
                )

                Alert.objects.create(
                    case=case,
                    alert_type="STATUS_UPDATE",
                    message=(
                        f"Case {case.case_number} status "
                        f"changed from "
                        f"{old_status_display} "
                        f"to {case.get_status_display()}."
                    )
                )

                messages.success(
                    request,
                    f"Case {case.case_number} status "
                    "updated successfully."
                )

            else:

                messages.info(
                    request,
                    "The case status is already set to "
                    f"{case.get_status_display()}."
                )

        else:

            messages.error(
                request,
                "Invalid case status selected."
            )

    return redirect(
        "case_detail",
        case_id=case.id
    )


# ALERTS
@login_required(login_url="login")
def alerts(request):

    alert_list = Alert.objects.all().order_by(
        "-created_at"
    )

    return render(
        request,
        "accounts/alerts.html",
        {
            "alerts": alert_list
        }
    )


# MARK SINGLE ALERT AS READ
@login_required(login_url="login")
def mark_alert_read(request, alert_id):

    alert = get_object_or_404(
        Alert,
        id=alert_id
    )

    if request.method == "POST":

        alert.is_read = True

        alert.save()

    return redirect(
        "alerts"
    )


# MARK ALL ALERTS AS READ
@login_required(login_url="login")
def mark_all_alerts_read(request):

    if request.method == "POST":

        Alert.objects.filter(
            is_read=False
        ).update(
            is_read=True
        )

        messages.success(
            request,
            "All alerts have been marked as read."
        )

    return redirect(
        "alerts"
    )


# REPORTS
@login_required(login_url="login")
def reports(request):

    total_cases = MissingChild.objects.count()

    reported_cases = MissingChild.objects.filter(
        status="REPORTED"
    ).count()

    active_cases = MissingChild.objects.filter(
        status="ACTIVE"
    ).count()

    under_review = MissingChild.objects.filter(
        status="UNDER_REVIEW"
    ).count()

    resolved_cases = MissingChild.objects.filter(
        status="RESOLVED"
    ).count()

    recent_cases = MissingChild.objects.order_by(
        "-created_at"
    )[:10]

    total_alerts = Alert.objects.count()

    reported_percentage = 0
    active_percentage = 0
    under_review_percentage = 0
    resolved_percentage = 0

    if total_cases > 0:

        reported_percentage = (
            reported_cases / total_cases
        ) * 100

        active_percentage = (
            active_cases / total_cases
        ) * 100

        under_review_percentage = (
            under_review / total_cases
        ) * 100

        resolved_percentage = (
            resolved_cases / total_cases
        ) * 100

    context = {

        "total_cases": total_cases,

        "reported_cases": reported_cases,

        "active_cases": active_cases,

        "under_review": under_review,

        "resolved_cases": resolved_cases,

        "recent_cases": recent_cases,

        "total_alerts": total_alerts,

        "reported_percentage": round(
            reported_percentage,
            2
        ),

        "active_percentage": round(
            active_percentage,
            2
        ),

        "under_review_percentage": round(
            under_review_percentage,
            2
        ),

        "resolved_percentage": round(
            resolved_percentage,
            2
        ),

    }

    return render(
        request,
        "accounts/reports.html",
        context
    )


# DOWNLOAD EXCEL CASE REPORT
@login_required(login_url="login")
def download_report(request):

    cases = MissingChild.objects.all().order_by(
        "-created_at"
    )

    workbook = Workbook()

    worksheet = workbook.active

    worksheet.title = (
        "Missing Child Reports"
    )

    headers = [
        "Case Number",
        "Child Name",
        "Age",
        "Last Seen Location",
        "Status",
        "Registered Date",
        "Reported By",
        "Child Photo"
    ]

    worksheet.append(
        headers
    )

    # Make header bold
    for cell in worksheet[1]:

        cell.font = cell.font.copy(
            bold=True
        )

    # Column widths
    worksheet.column_dimensions["A"].width = 18
    worksheet.column_dimensions["B"].width = 25
    worksheet.column_dimensions["C"].width = 10
    worksheet.column_dimensions["D"].width = 30
    worksheet.column_dimensions["E"].width = 18
    worksheet.column_dimensions["F"].width = 18
    worksheet.column_dimensions["G"].width = 22
    worksheet.column_dimensions["H"].width = 20

    row = 2

    for case in cases:

        worksheet.cell(
            row=row,
            column=1,
            value=case.case_number
        )

        worksheet.cell(
            row=row,
            column=2,
            value=case.child_name
        )

        worksheet.cell(
            row=row,
            column=3,
            value=case.age
        )

        worksheet.cell(
            row=row,
            column=4,
            value=case.location_last_seen
        )

        worksheet.cell(
            row=row,
            column=5,
            value=case.get_status_display()
        )

        if case.created_at:

            worksheet.cell(
                row=row,
                column=6,
                value=case.created_at.strftime(
                    "%d-%m-%Y"
                )
            )

        if case.reported_by:

            worksheet.cell(
                row=row,
                column=7,
                value=case.reported_by.username
            )

        # Add child photo
        if case.photo:

            try:

                image = ExcelImage(
                    case.photo.path
                )

                image.width = 100
                image.height = 100

                worksheet.add_image(
                    image,
                    f"H{row}"
                )

                worksheet.row_dimensions[
                    row
                ].height = 85

            except Exception:

                worksheet.cell(
                    row=row,
                    column=8,
                    value="Photo unavailable"
                )

        else:

            worksheet.cell(
                row=row,
                column=8,
                value="No photo"
            )

        row += 1

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-"
            "officedocument.spreadsheetml.sheet"
        )
    )

    response["Content-Disposition"] = (
        'attachment; filename="'
        'safetrace_missing_child_reports.xlsx"'
    )

    workbook.save(
        response
    )

    return response


# EDIT CASE
@login_required(login_url="login")
def edit_case(request, case_id):

    case = get_object_or_404(
        MissingChild,
        id=case_id
    )

    profile = getattr(
        request.user,
        "userprofile",
        None
    )

    if profile is None or profile.role not in [
        "ADMIN",
        "POLICE"
    ]:

        messages.error(
            request,
            "You are not authorized to edit this case."
        )

        return redirect(
            "case_detail",
            case_id=case.id
        )

    if request.method == "POST":

        case.child_name = request.POST.get(
            "child_name"
        )

        case.age = request.POST.get(
            "age"
        )

        case.gender = request.POST.get(
            "gender"
        )

        case.date_last_seen = request.POST.get(
            "date_last_seen"
        )

        case.location_last_seen = request.POST.get(
            "location_last_seen"
        )

        case.description = request.POST.get(
            "description"
        )

        case.guardian_name = request.POST.get(
            "guardian_name"
        )

        case.guardian_contact = request.POST.get(
            "guardian_contact"
        )

        if request.FILES.get("photo"):

            case.photo = request.FILES.get(
                "photo"
            )

        case.save()

        messages.success(
            request,
            "Case details updated successfully."
        )

        return redirect(
            "case_detail",
            case_id=case.id
        )

    return render(
        request,
        "accounts/edit_case.html",
        {
            "case": case,
            "profile": profile
        }
    )


# DELETE CASE
@login_required(login_url="login")
def delete_case(request, case_id):

    case = get_object_or_404(
        MissingChild,
        id=case_id
    )

    profile = getattr(
        request.user,
        "userprofile",
        None
    )

    if profile is None or profile.role not in [
        "ADMIN",
        "POLICE"
    ]:

        messages.error(
            request,
            "You are not authorized to delete this case."
        )

        return redirect(
            "case_detail",
            case_id=case.id
        )

    if request.method == "POST":

        case_number = case.case_number
        child_name = case.child_name

        case.delete()

        messages.success(
            request,
            f"Case {case_number} for "
            f"{child_name} has been "
            "deleted successfully."
        )

        return redirect(
            "missing_cases"
        )

    return render(
        request,
        "accounts/delete_case.html",
        {
            "case": case
        }
    )


# LOGOUT
def logout_view(request):

    logout(request)

    return redirect(
        "home"
    )
def serve_media(request, path):
    """
    Serve uploaded media files for the SafeTrace prototype.
    """

    file_path = os.path.join(settings.MEDIA_ROOT, path)

    if not os.path.isfile(file_path):
        raise Http404("Media file not found")

    return FileResponse(
        open(file_path, "rb"),
        content_type="image/jpeg"
    )